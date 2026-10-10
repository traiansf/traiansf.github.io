"""Completează șablonul facultății cu datele din fisa_disciplinei.md.

Utilizare:
    python fisa_disciplinei.py [fisa.md] [sablon.doc|.docx] [iesire.docx]

Un șablon .doc este convertit întâi în .docx cu Microsoft Word (Windows).
Dacă Word este disponibil, rezultatul se exportă și în PDF, alături de .docx.
"""

import copy
import re
import subprocess
import sys
import tempfile
from pathlib import Path

import docx
from docx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
DEFAULT_MD = HERE / "fisa_disciplinei.md"
DEFAULT_TEMPLATE = HERE / "Fise_discipline" / "template-fisa-disciplinei-licenta-master-RO-2026-2027.doc"
DEFAULT_OUT = HERE / "fisa_disciplinei.docx"


# --- Markdown -------------------------------------------------------------

def parse_md(path):
    """Împarte fișa pe secțiuni: {"1": [...], "7.1": [...], ...}, fiecare cu
    tabele (liste de rânduri), liste numerotate și paragrafe, în ordine."""
    text = Path(path).read_text(encoding="utf-8")
    text = re.sub(r"^---\n.*?\n---\n", "", text, flags=re.S)
    text = re.sub(r"<!--.*?-->", "", text, flags=re.S)
    sections, key = {}, None
    for line in text.splitlines():
        m = re.match(r"#{2,3} (\d+(?:\.\d+)?)\.", line)
        if m:
            key = m.group(1)
            sections[key] = []
        elif key:
            sections[key].append(line)
    return {k: blocks(v) for k, v in sections.items()}


def blocks(lines):
    out, para, in_table = [], [], False

    def flush():
        if para:
            out.append(("p", " ".join(para)))
            para.clear()

    for line in lines:
        s = line.strip()
        if s.startswith("|"):
            flush()
            cells = [c.strip() for c in s.strip("|").split("|")]
            if all(re.fullmatch(r":?-+:?", c) for c in cells):
                continue
            if in_table:
                out[-1][1].append(cells)
            else:
                out.append(("table", [cells]))
                in_table = True
        elif re.match(r"\d+\. ", s):
            flush()
            if out and out[-1][0] == "list":
                out[-1][1].append(s)
            else:
                out.append(("list", [s]))
        elif not s:
            flush()
            in_table = False
        else:
            para.append(s)
    flush()
    return out


def tables(section):
    return [b[1] for b in section if b[0] == "table"]


def paragraphs(section):
    return [b[1] for b in section if b[0] == "p"]


def numbered(section):
    return [item for b in section if b[0] == "list" for item in b[1]]


# --- Word -----------------------------------------------------------------

INLINE = re.compile(r"(\*\*.+?\*\*|\*.+?\*|<https?://[^>]+>)")


def add_runs(par, text, rpr):
    for part in INLINE.split(text):
        if not part:
            continue
        bold = italic = False
        if part.startswith("**"):
            part, bold = part[2:-2], True
        elif part.startswith("*"):
            part, italic = part[1:-1], True
        elif part.startswith("<"):
            part = part[1:-1]
        run = par.add_run(part)
        if rpr is not None:
            new = copy.deepcopy(rpr)
            for tag in ("w:color", "w:b", "w:bCs", "w:i", "w:iCs"):
                for el in new.findall(qn(tag)):
                    new.remove(el)
            run._r.insert(0, new)
        run.bold = bold or None
        run.italic = italic or None


def base_rpr(par):
    for r in par._p.findall(qn("w:r")):
        if r.find(qn("w:rPr")) is not None:
            return r.find(qn("w:rPr"))
    ppr = par._p.find(qn("w:pPr"))
    return ppr.find(qn("w:rPr")) if ppr is not None else None


def set_cell(cell, texts, keep_first=False):
    """Înlocuiește conținutul celulei cu paragrafele date, în negru, păstrând
    fontul primului paragraf; cu keep_first, primul paragraf rămâne (ex. „Bibliografie:”)."""
    if isinstance(texts, str):
        texts = [texts]
    first = cell.paragraphs[0]
    rpr = copy.deepcopy(base_rpr(first))
    for p in cell.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    model = copy.deepcopy(first._p)
    for r in model.findall(qn("w:r")) + model.findall(qn("w:hyperlink")):
        model.remove(r)
    last = first._p
    if not keep_first:
        for r in first._p.findall(qn("w:r")) + first._p.findall(qn("w:hyperlink")):
            first._p.remove(r)
        if texts:
            add_runs(first, texts[0], rpr)
            texts = texts[1:]
    for t in texts:
        p = copy.deepcopy(model)
        last.addnext(p)
        last = p
        add_runs(docx.text.paragraph.Paragraph(p, cell), t, rpr)


def unmerge_vertical(row):
    for tc in row._tr.findall(qn("w:tc")):
        tcpr = tc.find(qn("w:tcPr"))
        if tcpr is not None:
            for vm in tcpr.findall(qn("w:vMerge")):
                tcpr.remove(vm)


def merge_first_column(rows):
    """Unește vertical prima coloană a rândurilor (aceeași etichetă)."""
    for i, row in enumerate(rows):
        tcpr = row._tr.findall(qn("w:tc"))[0].get_or_add_tcPr()
        vm = tcpr.makeelement(qn("w:vMerge"), {qn("w:val"): "restart"} if i == 0 else {})
        tcpr.append(vm)


def unique_cells(row):
    seen, out = set(), []
    for c in row.cells:
        if id(c._tc) not in seen:
            seen.add(id(c._tc))
            out.append(c)
    return out


def fill_pairs(table, md_rows):
    """Tabele cheie–valoare cu aceleași rânduri ca șablonul."""
    rows = [r for r in md_rows if any(r)]
    assert len(rows) == len(table.rows), (len(rows), len(table.rows))
    for row, (label, value) in zip(table.rows, rows):
        cells = unique_cells(row)
        code = re.match(r"\d\.\d\.", label)
        assert not code or cells[0].text.strip().startswith(code.group(0)), (label, cells[0].text)
        set_cell(cells[-1], value)


def fill_section3(table, sec):
    values = {}
    for t in tables(sec):
        for row in t:
            for i, c in enumerate(row[:-1]):
                m = re.search(r"(3\.\d)\.", c)
                if m and row[i + 1]:
                    values[m.group(1)] = row[i + 1]
    pos = {"3.1": (0, 1), "3.2": (0, 3), "3.3": (0, 5),
           "3.4": (1, 1), "3.5": (1, 3), "3.6": (1, 5),
           "3.7": (7, 1), "3.8": (8, 1), "3.9": (9, 1)}
    for code, (r, c) in pos.items():
        set_cell(table.rows[r].cells[c], values[code])
    study = [row for row in tables(sec)[1][1:]]
    for i, (label, hours) in enumerate(study):
        row = table.rows[3 + i]
        if row.cells[0].text.strip() != label:
            set_cell(row.cells[0], label)
        set_cell(row.cells[5], hours)


def fill_contents(table, md, sub):
    rows = list(table.rows)
    head = next(i for i, r in enumerate(rows) if r.cells[0].text.strip().startswith(sub + "."))
    bib = next(i for i in range(head + 1, len(rows)) if rows[i].cells[0].text.strip().startswith("Bibliografie"))
    content = rows[head + 1:bib]
    data = tables(md[sub])[0][1:]
    while len(content) < len(data):
        tr = copy.deepcopy(content[-1]._tr)
        content[-1]._tr.addnext(tr)
        content.append(docx.table._Row(tr, table))
    for extra in content[len(data):]:
        extra._tr.getparent().remove(extra._tr)
    for row in content[:len(data)]:
        unmerge_vertical(row)
    for row, cells in zip(content, data):
        for cell, text in zip(unique_cells(row), cells):
            set_cell(cell, text)
    refs = numbered(md[sub])
    if refs:
        set_cell(rows[bib].cells[0], refs, keep_first=True)


def fill_evaluation(table, sec):
    groups = {}
    for label, *vals in tables(sec)[0][1:]:
        groups.setdefault(label, []).append(vals)
    paras = paragraphs(sec)
    std = next(i for i, p in enumerate(paras) if p.startswith("**Standard minim"))
    first = True
    for row in list(table.rows):
        label = row.cells[0].text.strip()
        if label not in groups:
            continue
        rows = [row]
        for _ in groups[label][1:]:
            tr = copy.deepcopy(rows[-1]._tr)
            rows[-1]._tr.addnext(tr)
            rows.append(docx.table._Row(tr, table))
        if len(rows) > 1:
            merge_first_column(rows)
        for r, vals in zip(rows, groups[label]):
            cells = unique_cells(r)
            for j in range(3):
                set_cell(cells[1 + j], vals[j])
        if first:
            # restanța/mărirea se referă la examen: în rândul examenului
            set_cell(unique_cells(rows[0])[1], [groups[label][0][0]] + paras[:std])
            first = False
    std_row = next(i for i, r in enumerate(table.rows) if r.cells[0].text.strip().startswith("Standard minim"))
    set_cell(table.rows[std_row + 1].cells[0], paras[std + 1:])
    return dict(tables(sec)[1])


def fill_signatures(table, info):
    def replace(cell, pattern, value):
        for p in cell.paragraphs:
            if re.search(pattern, p.text):
                rpr = copy.deepcopy(base_rpr(p))
                for r in p._p.findall(qn("w:r")):
                    p._p.remove(r)
                add_runs(p, value, rpr)
                return
        raise ValueError(pattern)

    replace(table.rows[0].cells[0], r"\d\d\.\d\d\.\d{4}", info["Data completării"])
    replace(table.rows[1].cells[0], r"\d\d\.\d\d\.\d{4}", info["Data avizării în departament"])
    replace(table.rows[0].cells[2], r"\.{5,}", info["Titularul de curs"])
    replace(table.rows[1].cells[1], r"Prof\.|Conf\.", info["Director de departament"])


def word(src, dst, fmt):
    """Salvează src în dst cu Word (fmt: 16 = docx, 17 = pdf)."""
    ps = (f"$w = New-Object -ComObject Word.Application; $w.Visible = $false; "
          f"$d = $w.Documents.Open('{src}', $false, $true); $d.SaveAs2('{dst}', {fmt}); "
          f"$d.Close($false); $w.Quit()")
    subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True)


def main(md_path=DEFAULT_MD, template=DEFAULT_TEMPLATE, out=DEFAULT_OUT):
    md, template, out = parse_md(md_path), Path(template).resolve(), Path(out).resolve()
    if template.suffix.lower() == ".doc":
        converted = Path(tempfile.mkdtemp()) / (template.stem + ".docx")
        word(template, converted, 16)
        template = converted
    doc = docx.Document(template)
    t = doc.tables

    fill_pairs(t[0], tables(md["1"])[0])
    sec2 = dict((r[0], r[1]) for r in tables(md["2"])[0] if r[0])
    get = lambda code: next(v for k, v in sec2.items() if k.startswith(code))
    set_cell(t[1].rows[0].cells[2], get("2.1."))
    set_cell(t[1].rows[1].cells[5], get("2.2."))
    set_cell(t[1].rows[2].cells[5], get("2.3."))
    for code, col in (("2.4.", 1), ("2.5.", 3), ("2.6.", 6), ("2.7.", 8)):
        set_cell(t[1].rows[3].cells[col], get(code))
    fill_section3(t[2], md["3"])
    fill_pairs(t[3], tables(md["4"])[0])
    fill_pairs(t[4], tables(md["5"])[0])
    fill_pairs(t[5], tables(md["6"])[0])
    for sub in ("7.1", "7.2", "7.3", "7.4"):
        fill_contents(t[6], md, sub)
    set_cell(t[7].rows[0].cells[0], paragraphs(md["8"]))
    fill_signatures(t[9], fill_evaluation(t[8], md["9"]))

    doc.save(out)
    print("scris:", out)
    try:
        word(out, out.with_suffix(".pdf"), 17)
        print("scris:", out.with_suffix(".pdf"))
    except (OSError, subprocess.CalledProcessError) as e:
        print("PDF negenerat (Word indisponibil):", e)


if __name__ == "__main__":
    main(*sys.argv[1:])
